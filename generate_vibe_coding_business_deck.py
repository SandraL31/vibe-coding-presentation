from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor


# ---------- Theme ----------
NAVY = RGBColor(17, 24, 39)
BLUE = RGBColor(37, 99, 235)
LIGHT = RGBColor(245, 247, 250)
TEXT = RGBColor(31, 41, 55)
GRAY = RGBColor(107, 114, 128)
WHITE = RGBColor(255, 255, 255)
GREEN = RGBColor(22, 163, 74)
ORANGE = RGBColor(249, 115, 22)
RED = RGBColor(220, 38, 38)


def add_title(slide, title, subtitle=None):
    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(11.8), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.bold = True
    p.font.size = Pt(26)
    p.font.color.rgb = NAVY

    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.05), Inches(11.8), Inches(0.5))
        tf2 = sub_box.text_frame
        p2 = tf2.paragraphs[0]
        p2.text = subtitle
        p2.font.size = Pt(12)
        p2.font.color.rgb = GRAY


def add_footer(slide, text):
    footer = slide.shapes.add_textbox(Inches(0.7), Inches(7.0), Inches(11.6), Inches(0.35))
    tf = footer.text_frame
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.RIGHT
    p.font.size = Pt(9)
    p.font.color.rgb = GRAY


def add_bullet_box(slide, x, y, w, h, bullets, bullet_color=TEXT, font_size=18):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0
        p.bullet = True
        p.font.size = Pt(font_size)
        p.font.color.rgb = bullet_color
        p.space_after = Pt(8)


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Slide 1: Cover
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = LIGHT

# top accent bar
bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(0.18))
bar.fill.solid()
bar.fill.fore_color.rgb = BLUE
bar.line.fill.background()

# Brand / title
label = slide.shapes.add_textbox(Inches(0.8), Inches(0.55), Inches(5.0), Inches(0.5))
tf = label.text_frame
p = tf.paragraphs[0]
p.text = "AI 產品 / 轉型策略簡報"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = BLUE

title = slide.shapes.add_textbox(Inches(0.8), Inches(1.15), Inches(9.5), Inches(1.2))
tf2 = title.text_frame
p2 = tf2.paragraphs[0]
p2.text = "Vibe Coding"
p2.font.bold = True
p2.font.size = Pt(30)
p2.font.color.rgb = NAVY

subtitle = slide.shapes.add_textbox(Inches(0.8), Inches(2.1), Inches(8.8), Inches(1.0))
tf3 = subtitle.text_frame
p3 = tf3.paragraphs[0]
p3.text = "AI 協作的新型程式開發模式\n將自然語言轉化為產品原型與可執行解決方案"
p3.font.size = Pt(18)
p3.font.color.rgb = TEXT

# right side block
card = slide.shapes.add_shape(1, Inches(9.1), Inches(1.3), Inches(3.2), Inches(2.6))
card.fill.solid()
card.fill.fore_color.rgb = NAVY
card.line.fill.background()

card_tf = card.text_frame
card_tf.word_wrap = True
card_tf.margin_left = 0.2
card_tf.margin_right = 0.2
card_tf.margin_top = 0.2
card_tf.margin_bottom = 0.2
for i, line in enumerate(["AI 協作", "自然語言需求", "快速原型", "產品迭代"]):
    p = card_tf.paragraphs[0] if i == 0 else card_tf.add_paragraph()
    p.text = line
    p.level = 0
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.space_after = Pt(8)

add_footer(slide, "資料來源：維基百科、AI 開發趨勢與產業討論 | 2026")

# Slide 2: What is Vibe Coding?
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = LIGHT
add_title(slide, "什麼是 Vibe Coding？", "一種以大型語言模型為核心的 AI 協作開發方式")

left_box = slide.shapes.add_shape(1, Inches(0.8), Inches(1.6), Inches(5.8), Inches(4.8))
left_box.fill.solid()
left_box.fill.fore_color.rgb = WHITE
left_box.line.color.rgb = BLUE
left_box.line.width = Pt(1.5)
left_tf = left_box.text_frame
left_tf.word_wrap = True
for i, line in enumerate([
    "• 不是傳統逐行手寫程式碼",
    "• 透過自然語言描述需求",
    "• AI 生成初版程式碼與功能",
    "• 由開發者進行測試、修正與迭代",
    "• 目標是快速驗證想法與原型",
    "• 重心從『寫程式』轉到『定義需求與協作』"
]):
    p = left_tf.paragraphs[0] if i == 0 else left_tf.add_paragraph()
    p.text = line
    p.level = 0
    p.font.size = Pt(20)
    p.font.color.rgb = TEXT
    p.bullet = True
    p.space_after = Pt(6)

right_box = slide.shapes.add_shape(1, Inches(7.2), Inches(1.8), Inches(5.2), Inches(3.6))
right_box.fill.solid()
right_box.fill.fore_color.rgb = BLUE
right_box.line.fill.background()
rtf = right_box.text_frame
rtf.word_wrap = True
for i, line in enumerate(["AI 需求輸入", "系統生成草案", "測試修正", "持續進化"]):
    p = rtf.paragraphs[0] if i == 0 else rtf.add_paragraph()
    p.text = line
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.space_after = Pt(12)

add_footer(slide, "關鍵說明：這是一種更偏 AI 協作與快速驗證的開發方法，而非完全替代工程實踐")

# Slide 3: business value
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = LIGHT
add_title(slide, "為什麼它值得企業注意？", "從原型效率、員工參與度與資源效率三個維度看價值")

# Three KPI-ish cards
cards = [
    (Inches(0.8), Inches(1.7), Inches(3.5), Inches(3.8), "效率提升", ["• 快速建立 MVP", "• 減少前期開發等待", "• 讓想法更快驗證"]),
    (Inches(4.7), Inches(1.7), Inches(3.8), Inches(3.8), "組織協作", ["• 降低非工程背景門檻", "• 行銷 / 業務可直接提出需求", "• 跨部門協作更順暢"]),
    (Inches(8.9), Inches(1.7), Inches(3.6), Inches(3.8), "成本優化", ["• 小型工具可快速落地", "• 更少重複開發", "• 更適合探索性專案"]),
]

for x, y, w, h, title, bullets in cards:
    c = slide.shapes.add_shape(1, x, y, w, h)
    c.fill.solid()
    c.fill.fore_color.rgb = WHITE
    c.line.color.rgb = BLUE
    c.line.width = Pt(1.3)
    tf = c.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.bold = True
    p.font.size = Pt(18)
    p.font.color.rgb = BLUE
    for line in bullets:
        p = tf.add_paragraph()
        p.text = line
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT
        p.bullet = True
        p.space_after = Pt(8)

add_footer(slide, "商業價值：將 AI 從『工具』轉變為『開發輔助與創新加速器』")

# Slide 4: risks & recommendations
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = LIGHT
add_title(slide, "風險與建議策略", "善用 AI 協作，但不能忽略工程治理與品質控制")

left = slide.shapes.add_shape(1, Inches(0.8), Inches(1.8), Inches(5.7), Inches(4.6))
left.fill.solid()
left.fill.fore_color.rgb = WHITE
left.line.color.rgb = RED
left.line.width = Pt(1.3)
ltf = left.text_frame
ltf.word_wrap = True
for i, line in enumerate([
    "⚠  代碼責任難以界定",
    "⚠  維護與擴充成本可能增加",
    "⚠  安全性與合規風險需審慎處理",
    "⚠  複雜專案仍依賴工程基礎能力",
    "⚠  AI 產出可能偏向「能動」但不一定「穩定」"
]):
    p = ltf.paragraphs[0] if i == 0 else ltf.add_paragraph()
    p.text = line
    p.font.size = Pt(18)
    p.font.color.rgb = TEXT
    p.bullet = True
    p.space_after = Pt(8)

right = slide.shapes.add_shape(1, Inches(7.0), Inches(1.8), Inches(5.5), Inches(4.6))
right.fill.solid()
right.fill.fore_color.rgb = NAVY
right.line.fill.background()
rtf = right.text_frame
for i, line in enumerate([
    "✅ 建立 AI 開發規範",
    "✅ 著重需求與測試，而非只看生成結果",
    "✅ 保留工程審查與版本控制",
    "✅ 將 AI 用於原型與加速，而非完全取代團隊判斷",
    "✅ 以低風險場景先行，再擴展到核心業務流程"
]):
    p = rtf.paragraphs[0] if i == 0 else rtf.add_paragraph()
    p.text = line
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.bullet = True
    p.space_after = Pt(10)

add_footer(slide, "結論：Vibe Coding 是代表 AI 導向開發的加速器，不是取代工程思維的替代品")

prs.save("vibe_coding_business_deck.pptx")
print("PPTX generated: vibe_coding_business_deck.pptx")
